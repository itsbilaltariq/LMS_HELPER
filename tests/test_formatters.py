"""Unit tests for multi-format deliverable generators (ipynb, docx, xlsx)."""

import json
from pathlib import Path

from docx import Document
from openpyxl import load_workbook

from workers.formatters import (
    build_notebook_structure,
    convert_solution_to_notebook,
    convert_solution_to_word,
    create_excel_spreadsheet,
    create_jupyter_notebook,
    create_word_document,
)


def test_build_notebook_structure():
    """Verify Jupyter notebook schema adheres to nbformat v4."""
    cells = [
        {"cell_type": "markdown", "source": "# Lab 1 Title\nIntroduction"},
        {"cell_type": "code", "source": "def solution():\n    return 42"},
    ]
    nb = build_notebook_structure(cells)

    assert nb["nbformat"] == 4
    assert nb["nbformat_minor"] == 5
    assert len(nb["cells"]) == 2
    assert nb["cells"][0]["cell_type"] == "markdown"
    assert nb["cells"][1]["cell_type"] == "code"
    assert nb["metadata"]["language_info"]["name"] == "python"


def test_create_jupyter_notebook(tmp_path: Path):
    """Verify creation and JSON persistence of .ipynb files."""
    out_file = tmp_path / "lab_submission.ipynb"
    cells = [
        {"cell_type": "markdown", "source": "# Lab Exercise"},
        {"cell_type": "code", "source": "print('NUST')"},
    ]
    result_path = create_jupyter_notebook(cells, out_file)

    assert result_path.exists()
    content = json.loads(result_path.read_text(encoding="utf-8"))
    assert content["nbformat"] == 4
    assert len(content["cells"]) == 2


def test_convert_solution_to_notebook(tmp_path: Path):
    """Verify high-level conversion of solution code into submission notebook."""
    out_file = tmp_path / "complete_lab.ipynb"
    result = convert_solution_to_notebook(
        title="CS212 BST Lab",
        solution_code="class BSTNode: ...",
        walkthrough="This algorithm achieves O(log N) average time.",
        output_path=out_file,
        viva_questions=["What is tree rotation?"],
    )

    assert result.exists()
    data = json.loads(result.read_text(encoding="utf-8"))
    assert len(data["cells"]) == 4  # intro, code header, code, viva
    # Check that code cell has the solution code
    code_cells = [c for c in data["cells"] if c["cell_type"] == "code"]
    assert len(code_cells) == 1
    assert "class BSTNode: ..." in "".join(code_cells[0]["source"])


def test_create_word_document(tmp_path: Path):
    """Verify Word document (.docx) creation and content."""
    out_docx = tmp_path / "lab_report.docx"
    sections = [
        {"heading": "Abstract", "content": "Analysis of sorting algorithms."},
        {"heading": "Code", "content": "int main() { return 0; }", "is_code": True},
    ]
    created = create_word_document("Sorting Lab", sections, out_docx)

    assert created.exists()
    doc = Document(str(created))
    paras = [p.text for p in doc.paragraphs if p.text.strip()]
    assert "Sorting Lab" in paras
    assert "Analysis of sorting algorithms." in paras
    assert "int main() { return 0; }" in paras


def test_convert_solution_to_word(tmp_path: Path):
    """Verify conversion of solution code into formatted Word document."""
    out_docx = tmp_path / "submission_report.docx"
    created = convert_solution_to_word(
        title="DLD Lab 4",
        solution_code="module adder(); endmodule",
        walkthrough="Full adder with gate level modeling.",
        output_path=out_docx,
        viva_questions=["What is propagation delay?"],
    )

    assert created.exists()
    doc = Document(str(created))
    full_text = "\n".join([p.text for p in doc.paragraphs])
    assert "DLD Lab 4" in full_text
    assert "Full adder with gate level modeling." in full_text
    assert "module adder(); endmodule" in full_text
    assert "What is propagation delay?" in full_text


def test_create_excel_spreadsheet(tmp_path: Path):
    """Verify Excel workbook (.xlsx) creation with rows and sheets."""
    out_xlsx = tmp_path / "results.xlsx"
    sheets = {
        "TestCases": [
            ["Test ID", "Input", "Expected", "Result"],
            [1, "arr=[3,1,2]", "[1,2,3]", "PASS"],
            [2, "arr=[]", "[]", "PASS"],
        ],
        "Metrics": [
            ["Metric", "Value"],
            ["Time Complexity", "O(N log N)"],
            ["Space Complexity", "O(1)"],
        ],
    }
    created = create_excel_spreadsheet(sheets, out_xlsx)

    assert created.exists()
    wb = load_workbook(str(created))
    assert "TestCases" in wb.sheetnames
    assert "Metrics" in wb.sheetnames
    ws_tests = wb["TestCases"]
    assert ws_tests.cell(row=2, column=4).value == "PASS"
    ws_metrics = wb["Metrics"]
    assert ws_metrics.cell(row=2, column=2).value == "O(N log N)"
